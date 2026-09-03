import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import OrderedDict

SRC = r'C:\Users\PERUVIAN FIT\Desktop\SEGUIMIENTOSMS\MISHA SMS 1 SE.xlsx'
OUT = r'C:\Users\PERUVIAN FIT\Desktop\SEGUIMIENTOSMS\REPORTE SMS-SUMMER27.xlsx'

SIZES = ['one size','4-8y','0-6m','6-12m','3-6m','6-9m','9-12m','2y','3y','4Y','5y','6Y','8y','10y']
# indices de columna (0-based) segun estructura
PED_IDX = list(range(3,17))   # D..Q
PED_TOTAL = 17                # R
PROG_OFFSET = 15              # S = 18 -> D(3)+15
PROG_IDX = [x+15 for x in PED_IDX]  # S..AF
PROG_TOTAL = 32               # AG
IDX_NAME = 33                 # AH
IDX_TELA = 34                 # AI
IDX_COLOR = 35                # AJ
IDX_STATUS = 36               # AK
IDX_INGRESO = 37              # AL
IDX_OBS = 38                  # AM

wb_src = openpyxl.load_workbook(SRC, data_only=True)
ws = wb_src['STATUS SMS-SUMMER27']

# --- Agrupar por Style y por Color ---
estilos = OrderedDict()
for row in ws.iter_rows(min_row=13, values_only=True):
    po = row[0]; styl = row[1]
    if not styl:
        continue
    color = str(row[IDX_COLOR]).strip() if row[IDX_COLOR] else ''
    fila = {
        'po': str(po).strip() if po else '',
        'style': str(styl).strip(),
        'tendido': row[2] if row[2] is not None else '',
        'pedido': [row[i] if row[i] is not None else 0 for i in PED_IDX],
        'pedido_total': row[PED_TOTAL] if row[PED_TOTAL] is not None else 0,
        'prog': [row[i] if row[i] is not None else 0 for i in PROG_IDX],
        'prog_total': row[PROG_TOTAL] if row[PROG_TOTAL] is not None else 0,
        'name': str(row[IDX_NAME]).strip() if row[IDX_NAME] else '',
        'tela': str(row[IDX_TELA]).strip() if row[IDX_TELA] else '',
        'color': color,
        'status': str(row[IDX_STATUS]).strip() if row[IDX_STATUS] else '',
        'ingreso': row[IDX_INGRESO] if row[IDX_INGRESO] is not None else '',
        'obs': str(row[IDX_OBS]).strip() if row[IDX_OBS] else '',
    }
    es = estilos.setdefault(fila['style'], {'nombre': fila['name'], 'tela': fila['tela'], 'colores': OrderedDict()})
    fila['name'] = es['nombre'] or fila['name']
    fila['tela'] = es['tela'] or fila['tela']
    es['colores'].setdefault(fila['color'] or '(sin color)', fila)

# --- Crear libro de salida ---
out = openpyxl.Workbook()
out.remove(out.active)

# estilos comunes
TITULO_FONT = Font(bold=True, size=13, color='FFFFFF')
TITULO_FILL = PatternFill('solid', fgColor='1F4E78')
HEAD_FONT = Font(bold=True, size=10, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='2E75B6')
SUBHEAD_FONT = Font(bold=True, size=10)
GRUPO_FONT = Font(bold=True, size=10, color='FFFFFF')
GRUPO_FILL = PatternFill('solid', fgColor='548235')
thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal='center', vertical='center')
LEFT = Alignment(horizontal='left', vertical='center')

# ---------- Hoja 1: RESUMEN GENERAL ----------
ws1 = out.create_sheet('RESUMEN GENERAL')
ws1['A1'] = 'REPORTE SMS-SUMMER27 - RESUMEN GENERAL'
ws1['A1'].font = TITULO_FONT; ws1['A1'].fill = TITULO_FILL
ws1.merge_cells('A1:H1')

headers = ['#','Style','Nombre (Prenda)','# Variantes (Colores)','Total PEDIDO','Total PROGRAMADO','Tela Principal','Estados de Tela']
for j,h in enumerate(headers, start=1):
    c = ws1.cell(3,j,h); c.font=HEAD_FONT; c.fill=HEAD_FILL; c.alignment=CENTER; c.border=BORDER

grand_pedido = 0
grand_prog = 0
row = 4
for n,(styl,es) in enumerate(estilos.items(), start=1):
    total_pedido = sum(v['pedido_total'] for v in es['colores'].values())
    total_prog = sum(v['prog_total'] for v in es['colores'].values())
    grand_pedido += total_pedido; grand_prog += total_prog
    estados = OrderedDict()
    for v in es['colores'].values():
        st = v['status'] or '(sin estado)'
        estados[st] = estados.get(st,0)+1
    estados_str = ', '.join(f'{k} ({v})' for k,v in estados.items())
    vals = [n, styl, es['nombre'], len(es['colores']), total_pedido, total_prog, es['tela'], estados_str]
    for j,val in enumerate(vals, start=1):
        c = ws1.cell(row,j,val); c.border=BORDER
        c.alignment = LEFT if j in (3,7,8) else CENTER
    row += 1

tot_row = row
ws1.cell(tot_row,1,'TOTAL').font = GRUPO_FONT; ws1.cell(tot_row,1).fill=GRUPO_FILL
ws1.cell(tot_row,2,'').fill=GRUPO_FILL
ws1.cell(tot_row,4,str(len(estilos))).font=GRUPO_FONT
ws1.cell(tot_row,5,grand_pedido).font=GRUPO_FONT
ws1.cell(tot_row,6,grand_prog).font=GRUPO_FONT
for col in range(1,9):
    ws1.cell(tot_row,col).fill=GRUPO_FILL
    ws1.cell(tot_row,col).border=BORDER
    ws1.cell(tot_row,col).alignment=LEFT if col==7 else CENTER

widths = [5,12,18,15,14,16,22,30]
for j,w in enumerate(widths, start=1):
    ws1.column_dimensions[get_column_letter(j)].width = w
ws1.freeze_panes = 'A4'

# ---------- Hoja 2: DETALLE POR ESTILO ----------
ws2 = out.create_sheet('DETALLE POR ESTILO')
ws2['A1'] = 'REPORTE SMS-SUMMER27 - DETALLE POR ESTILO (Style + Color Cuerpo)'
ws2['A1'].font = TITULO_FONT; ws2['A1'].fill = TITULO_FILL
ws2.merge_cells('A1:M1')

# headers detalle
det_headers = ['Style','Color Cuerpo','Nombre (Prenda)','Tela Principal','Total PEDIDO','Total PROGRAMADO','STATUS TELA','Ingreso Cost','Observaciones','Dif. (Prog-Ped)','Tendido','PO','']
for j,h in enumerate(det_headers, start=1):
    c = ws2.cell(3,j,h); c.font=HEAD_FONT; c.fill=HEAD_FILL; c.alignment=CENTER; c.border=BORDER

r = 4
for styl,es in estilos.items():
    ws2.cell(r,1,styl).font=GRUPO_FONT; ws2.cell(r,1).fill=GRUPO_FILL; ws2.cell(r,1).border=BORDER
    ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=13)
    for j in range(1,14):
        ws2.cell(r,j).fill=GRUPO_FILL; ws2.cell(r,j).border=BORDER
    r += 1
    for color,fila in es['colores'].items():
        dif = int(fila['prog_total'] or 0) - int(fila['pedido_total'] or 0)
        vals = [styl, fila['color'], fila['name'], fila['tela'],
                fila['pedido_total'], fila['prog_total'], fila['status'],
                fila['ingreso'] if isinstance(fila['ingreso'],str) else (fila['ingreso'].strftime('%d-%m-%Y') if hasattr(fila['ingreso'],'strftime') else fila['ingreso']),
                fila['obs'], dif]
        for j,val in enumerate(vals, start=1):
            c = ws2.cell(r,j,val); c.border=BORDER
            if j==2: c.alignment=LEFT
            elif j in (5,6,10): c.alignment=CENTER
            else: c.alignment=LEFT
        ws2.cell(r,11,fila['tendido'])
        ws2.cell(r,12,fila['po'])
        if dif<0:
            ws2.cell(r,10).fill = PatternFill('solid', fgColor='FFC7CE')
            ws2.cell(r,10).font = Font(color='9C0006')
        elif dif>0:
            ws2.cell(r,10).fill = PatternFill('solid', fgColor='C6EFCE')
            ws2.cell(r,10).font = Font(color='006100')
        r += 1
    r += 1  # fila en blanco entre estilos

det_widths = [12,22,18,20,12,15,14,14,38,12,10,12,6]
for j,w in enumerate(det_widths, start=1):
    ws2.column_dimensions[get_column_letter(j)].width = w
ws2.freeze_panes = 'A4'

out.save(OUT)
print('Reporte generado en:', OUT)
print('Estilos:', len(estilos), '| Variantes (colores):', sum(len(e['colores']) for e in estilos.values()))
print('Total PEDIDO:', grand_pedido, '| Total PROGRAMADO:', grand_prog)
